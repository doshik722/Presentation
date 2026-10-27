# Здесь проверяем, что сборщик создаёт открываемый PPTX и сообщает о неотрисованных элементах.
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

from presentation_rl.dsl import parse
from presentation_rl.rendering import THEMES, build_from_text, build_presentation, pick_theme

EXAMPLES = Path(__file__).resolve().parents[1] / "data" / "examples"


def read_example(name: str) -> str:
    return (EXAMPLES / name).read_text(encoding="utf-8")


def shape_texts(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize("name", ["gradient_descent.txt", "serf_reform.txt"])
def test_good_examples_build(tmp_path, name):
    deck, report = build_from_text(read_example(name), tmp_path / "deck.pptx")
    assert report.pptx_created
    assert report.slide_count == 12
    counts = report.status_counts()
    assert counts.get("render_error", 0) == 0
    assert counts.get("overflow", 0) == 0
    assert counts.get("text_overflow", 0) == 0
    # На каждом слайде есть визуальный элемент, и ни один слайд не пустой.
    assert not [s.slide for s in report.slides if s.sparse]
    assert all(s.visuals >= 1 for s in report.slides)
    # Картинки пока заглушки, и отчёт это отражает.
    assert counts["placeholder"] == sum(len(s.elements_of("IMAGE")) for s in deck.slides)

    prs = Presentation(tmp_path / "deck.pptx")
    for slide_data, slide in zip(deck.slides, prs.slides):
        text = shape_texts(slide)
        assert slide_data.title in text
        for paragraph in slide_data.elements_of("BULLET", "TEXT"):
            assert paragraph.data["text"] in text
        if slide_data.notes:
            assert slide.notes_slide.notes_text_frame.text == "\n".join(slide_data.notes)


def test_font_sizes_are_fixed(tmp_path):
    long_text = "#TEXT: " + "слово " * 60
    _, report = build_from_text("#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: Заголовок\n#BULLET: пункт\n" + long_text, tmp_path / "f.pptx")
    prs = Presentation(tmp_path / "f.pptx")
    sizes = {
        run.text.strip()[:5]: run.font.size.pt
        for shape in prs.slides[1].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    }
    assert sizes["Загол"] == 30
    assert sizes["пункт"] == 20
    assert sizes["слово"] == 20


def test_text_has_no_bullets_but_bullets_do(tmp_path):
    build_from_text("#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: З\n#TEXT: абзац\n#BULLET: пункт", tmp_path / "b.pptx")
    prs = Presentation(tmp_path / "b.pptx")
    marks = {}
    for shape in prs.slides[1].shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            p_pr = paragraph._p.pPr
            marks[paragraph.text] = p_pr is not None and p_pr.find(qn("a:buChar")) is not None
    assert marks["абзац"] is False
    assert marks["пункт"] is True


def test_charts_are_native(tmp_path):
    deck, _ = build_from_text(read_example("serf_reform.txt"), tmp_path / "deck.pptx")
    prs = Presentation(tmp_path / "deck.pptx")
    charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    expected = sum(len(s.elements_of("CHART")) for s in deck.slides)
    assert len(charts) == expected == 2
    assert list(charts[0].plots[0].categories) == ["помещичьи", "государственные", "удельные"]


def test_overflow_is_reported(tmp_path):
    _, report = build_from_text(read_example("broken.txt"), tmp_path / "broken.pptx")
    assert report.pptx_created and report.slide_count == 2
    overflow = [item for item in report.items if item.status == "overflow"]
    assert [(item.kind, item.line) for item in overflow] == [("IMAGE", 18)]


# Формулы, которые должны рисоваться. Пополняем из реальных ответов моделей.
GOOD_FORMULAS = [
    r"\theta_{t+1} = \theta_t - \eta \nabla_\theta J(\theta_t)",
    r"\hat{m}_t = \frac{m_t}{1 - \beta_1^t}",  # Adam, T-lite, 29.09
    r"6\text{CO}_2 + 6\text{H}_2\text{O} \rightarrow \text{C}_6\text{H}_{12}\text{O}_6 + 6\text{O}_2",
    r"x_{\text{нов}} = x_{\text{стар}} - \eta f'(x)",  # русский текст внутри \text
    # правильный LaTeX, которого нет в mathtext: рисуется после замен (rendering/mathtext.py)
    r"6\text{CO}_2 + 6\text{H}_2\text{O} \xrightarrow{\text{свет}} \text{C}_6\text{H}_{12}\text{O}_6 + 6\text{O}_2",  # T-lite, 29.09
    r"\mathrm{ReLU}(x) = \begin{cases} x & x > 0 \\ 0 & x \le 0 \end{cases}",
    r"\begin{aligned} x_{k+1} &= x_k - 2\eta x_k \\ &= x_k (1 - 2\eta) \end{aligned}",
    r"\displaystyle \sum\limits_{i=1}^{n} \big( y_i - \hat{y}_i \big)^2 \tag{1}",
]

# Формулы, которые должны получить render_error.
BAD_FORMULAS = [
    r"\frac{1}{",  # незакрытая скобка: matplotlib падает
    r"\frac{d}{dx}(x^\text{\textasciicircum}2) = 2x",  # T-lite, 29.09: рисуется с заглушкой вместо символа
    # правильный LaTeX, который пока не переписываем: модель учится его избегать
    r"H = \begin{pmatrix} a & b \\ c & d \end{pmatrix}",
    r"\ce{6CO2 + 6H2O -> C6H12O6 + 6O2}",
]


def formula_statuses(tmp_path, latex):
    _, report = build_from_text(f"#SLIDE\n#TITLE: Т\n#FORMULA: {latex}", tmp_path / "f.pptx")
    assert report.pptx_created  # сломанная формула не роняет сборку
    return [(i.kind, i.status) for i in report.items]


@pytest.mark.parametrize("latex", GOOD_FORMULAS)
def test_good_formula_is_ok(tmp_path, latex):
    assert formula_statuses(tmp_path, latex) == [("FORMULA", "ok")]


@pytest.mark.parametrize("latex", BAD_FORMULAS)
def test_bad_formula_is_render_error(tmp_path, latex):
    assert formula_statuses(tmp_path, latex) == [("FORMULA", "render_error")]


def test_plot_render_error_is_reported(tmp_path):
    _, report = build_from_text("#SLIDE\n#TITLE: Т\n#PLOT: y = log(x); x=-5..-1", tmp_path / "p.pptx")
    assert [(i.kind, i.status) for i in report.items] == [("PLOT", "render_error")]


@pytest.mark.parametrize("plot, status", [
    ("y = 3x + 1; x=0..10", "ok"),  # честная прямая — не ошибка
    ("y = x^2; x=-2..2", "ok"),
    ("y = sin(x); x=0..0.1", "flat_plot"),  # нелинейная, но на отрезке выглядит прямой
    ("y = x^2; x=100..101", "flat_plot"),
    ("y = exp(x); x=0..100", "flat_plot"),  # выброс сжал кривую в линию
    ("y = log(x); x=-10..1", "undefined_plot"),  # почти везде не определена: рисуется, но это ошибка
])
def test_plot_quality_status(tmp_path, plot, status):
    _, report = build_from_text(f"#SLIDE\n#TITLE: Т\n#PLOT: {plot}", tmp_path / "p.pptx")
    assert [(i.kind, i.status) for i in report.items] == [("PLOT", status)]


def test_constant_line_chart_is_flat(tmp_path):
    text = "#SLIDE\n#TITLE: Т\n#CHART: type=line; title=Доли; x=2020,2021,2022; y=5,5,5"
    _, report = build_from_text(text, tmp_path / "c.pptx")
    assert [(i.kind, i.status) for i in report.items] == [("CHART", "flat_plot")]


def test_formula_limit(tmp_path):
    text = "#SLIDE\n#TITLE: Т\n" + "\n".join(f"#FORMULA: x_{i}" for i in range(4))
    _, report = build_from_text(text, tmp_path / "f.pptx")
    assert report.status_counts() == {"ok": 3, "overflow": 1}
    assert [i.line for i in report.items if i.status == "overflow"] == [6]


def test_empty_deck_creates_no_file(tmp_path):
    report = build_presentation(parse("просто текст"), tmp_path / "none.pptx")
    assert not report.pptx_created
    assert not (tmp_path / "none.pptx").exists()


def test_text_overflow_is_reported_not_shrunk(tmp_path):
    long_bullets = "\n".join("#BULLET: " + "очень длинный пункт " * 30 for _ in range(8))
    _, report = build_from_text("#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: З\n" + long_bullets, tmp_path / "t.pptx")
    assert report.count("text_overflow") == 1
    prs = Presentation(tmp_path / "t.pptx")
    runs = [r for s in prs.slides[1].shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
    assert {r.font.size.pt for r in runs if "пункт" in r.text} == {20}


def test_section_slide_with_title_only(tmp_path):
    _, report = build_from_text("#SLIDE\n#TITLE: Начало\n#SLIDE\n#TITLE: Раздел", tmp_path / "s.pptx")
    assert report.pptx_created and report.slide_count == 2
    prs = Presentation(tmp_path / "s.pptx")
    assert "Раздел" in shape_texts(prs.slides[1])


def test_sparse_slide_is_detected(tmp_path):
    text = "#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: Мало\n#TEXT: Одна короткая фраза.\n#SLIDE\n#TITLE: С графиком\n#TEXT: Та же фраза.\n#PLOT: y = x^2; x=-1..1"
    _, report = build_from_text(text, tmp_path / "s.pptx")
    stats = {s.slide: s for s in report.slides}
    assert stats[2].sparse and stats[2].visuals == 0
    assert not stats[3].sparse and stats[3].visuals == 1


def test_bullets_and_text_share_one_text_box(tmp_path):
    build_from_text("#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: З\n#BULLET: пункт\n#TEXT: абзац", tmp_path / "o.pptx")
    prs = Presentation(tmp_path / "o.pptx")
    boxes = [s for s in prs.slides[1].shapes if s.has_text_frame and "пункт" in s.text_frame.text]
    assert len(boxes) == 1 and "абзац" in boxes[0].text_frame.text


def _contrast(a: str, b: str) -> float:
    def lum(h: str) -> float:
        c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    hi, lo = sorted([lum(a), lum(b)], reverse=True)
    return (hi + 0.05) / (lo + 0.05)


@pytest.mark.parametrize("name", sorted(THEMES))
def test_theme_text_is_readable(name):
    # Пороги WCAG: основной текст — 7, мелкие подписи — 4,5, белый текст на акцентном узле — 3 (жирный).
    t = THEMES[name]
    assert _contrast(t.ink, "FFFFFF") >= 7 and _contrast(t.body, "FFFFFF") >= 7
    assert _contrast(t.muted, "FFFFFF") >= 4.5 and _contrast(t.muted, t.accent_soft) >= 4.5
    assert _contrast(t.ink, t.accent_soft) >= 7 and _contrast(t.accent, "FFFFFF") >= 3
    assert _contrast(t.title_text, t.title_bg) >= 7 and _contrast(t.title_body, t.title_bg) >= 4.5


@pytest.mark.parametrize("name", sorted(THEMES))
def test_every_theme_builds(tmp_path, name):
    _, report = build_from_text(read_example("gradient_descent.txt"), tmp_path / "deck.pptx", theme=THEMES[name])
    assert report.pptx_created and report.status_counts().get("render_error", 0) == 0


def test_auto_theme_is_stable_and_varied():
    assert pick_theme("auto", "Фотоэффект") == pick_theme("auto", "Фотоэффект")
    assert len({pick_theme("auto", f"тема {i}") for i in range(60)}) == len(THEMES)
    assert pick_theme("forest", "что угодно") == THEMES["forest"]
    with pytest.raises(ValueError):
        pick_theme("rainbow")


def test_long_title_does_not_cover_subtitle(tmp_path):
    text = "#SLIDE\n#TITLE: Нормативно-правовое регулирование безопасности\n#TEXT: Подзаголовок\n#IMAGE: кодексы"
    build_from_text(text, tmp_path / "deck.pptx")
    shapes = [s for s in Presentation(tmp_path / "deck.pptx").slides[0].shapes if s.has_text_frame]
    title = next(s for s in shapes if "Нормативно" in s.text_frame.text)
    subtitle = next(s for s in shapes if s.text_frame.text == "Подзаголовок")
    assert title.text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM
    assert title.top + title.height <= subtitle.top
    # Заголовок в колонке 7 дюймов при 40 pt: в PowerPoint — 3 строки (замер 27.10.2026, sft_t018_gemini).
    assert title.height >= Pt(40 * 1.2 * 3)


SOURCES = {
    "d1": {"title": "Градиентный спуск", "section": "", "url": "https://ru.wikipedia.org/wiki/A", "lang": "ru"},
    "d1#2": {"title": "Градиентный спуск", "section": "Выбор шага", "url": "https://ru.wikipedia.org/wiki/A#B", "lang": "ru"},
    "d3": {"title": "Gradient descent", "section": "", "url": "https://en.wikipedia.org/wiki/G", "lang": "en"},
    "d4": {"title": "Метод Ньютона", "section": "", "url": "https://ru.wikipedia.org/wiki/N", "lang": "ru"},
}


def _slide_texts(path):
    from pptx import Presentation

    return [" ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
            for slide in Presentation(path).slides]


def test_sources_line_numbers_and_sources_slide(tmp_path):
    text = """#SLIDE
#TITLE: Градиентный спуск
#SLIDE
#TITLE: Шаг
#BULLET: Шаг подбирают
#SOURCE: d1#2
#SLIDE
#TITLE: Сравнение
#BULLET: Три источника
#SOURCE: d1
#SOURCE: d3
#SOURCE: d4
#SOURCE: d1
#SOURCE: d9
"""
    path = tmp_path / "s.pptx"
    _, report = build_from_text(text, path, sources=SOURCES.get)
    texts = _slide_texts(path)
    assert "Источник: Градиентный спуск, раздел «Выбор шага» (Википедия)" in texts[1]
    assert "Источники: 2, 3, 4" in texts[2]
    assert report.sources_slide and report.slide_count == 4 and "Источники" in texts[3]
    assert "3. Gradient descent (англ. Википедия)" in texts[3]
    assert [s["number"] for s in report.sources] == [1, 2, 3, 4]
    assert report.count("bad_source") == 1 and report.count("ok") >= 4


def test_sources_without_resolver_are_errors(tmp_path):
    path = tmp_path / "s.pptx"
    _, report = build_from_text("#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: А\n#BULLET: б\n#SOURCE: d1\n", path)
    assert report.count("bad_source") == 1 and not report.sources_slide and report.slide_count == 2
    assert "Источник" not in " ".join(_slide_texts(path))



def test_hidden_sources_are_checked_but_not_drawn(tmp_path):
    path = tmp_path / "s.pptx"
    text = "#SLIDE\n#TITLE: Т\n#SLIDE\n#TITLE: А\n#BULLET: б\n#SOURCE: d1\n#SOURCE: d3\n#SOURCE: d4\n#SOURCE: d9\n"
    _, report = build_from_text(text, path, sources=SOURCES.get, show_sources=False)
    assert len(report.sources) == 3 and report.count("bad_source") == 1
    assert not report.sources_slide and report.slide_count == 2
    assert "Источник" not in " ".join(_slide_texts(path))
